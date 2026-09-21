from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from mi_llama.documents import chunk_document, detect_source_kind, parse_document
from mi_llama.domain import SourceKind


def test_detect_source_kind_supports_research_formats() -> None:
    assert detect_source_kind("paper.pdf", "application/octet-stream") is SourceKind.PDF
    assert detect_source_kind("book.epub", "application/octet-stream") is SourceKind.EPUB
    assert detect_source_kind("notes.md", "text/plain") is SourceKind.MARKDOWN
    assert detect_source_kind("page.html", "application/octet-stream") is SourceKind.HTML
    assert detect_source_kind("draft.txt", "text/plain") is SourceKind.TEXT


def test_html_extraction_drops_non_readable_content() -> None:
    document = parse_document(
        filename="source.html",
        media_type="text/html",
        content=b"<html><body><h1>Visible</h1><script>secret()</script><p>Evidence text.</p></body></html>",
        max_extracted_chars=10_000,
    )

    assert document.kind is SourceKind.HTML
    assert "Visible" in document.sections[0].text
    assert "Evidence text." in document.sections[0].text
    assert "secret()" not in document.sections[0].text


def test_docx_extraction_includes_tables() -> None:
    payload = BytesIO()
    document = Document()
    document.add_paragraph("Opening paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Claim"
    table.cell(0, 1).text = "Evidence"
    document.save(payload)

    parsed = parse_document(
        filename="research.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=payload.getvalue(),
        max_extracted_chars=10_000,
    )

    text = parsed.sections[0].text
    assert "Opening paragraph" in text
    assert "Claim | Evidence" in text


def test_epub_respects_spine_order() -> None:
    payload = BytesIO()
    with ZipFile(payload, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "META-INF/container.xml",
            """<?xml version='1.0'?>
            <container xmlns='urn:oasis:names:tc:opendocument:xmlns:container'>
              <rootfiles><rootfile full-path='OEBPS/content.opf'/></rootfiles>
            </container>""",
        )
        archive.writestr(
            "OEBPS/content.opf",
            """<package xmlns='http://www.idpf.org/2007/opf'>
              <manifest>
                <item id='second' href='second.xhtml' media-type='application/xhtml+xml'/>
                <item id='first' href='first.xhtml' media-type='application/xhtml+xml'/>
              </manifest>
              <spine><itemref idref='first'/><itemref idref='second'/></spine>
            </package>""",
        )
        archive.writestr("OEBPS/first.xhtml", "<html><body><p>First chapter</p></body></html>")
        archive.writestr("OEBPS/second.xhtml", "<html><body><p>Second chapter</p></body></html>")

    parsed = parse_document(
        filename="book.epub",
        media_type="application/epub+zip",
        content=payload.getvalue(),
        max_extracted_chars=10_000,
    )

    assert [section.text for section in parsed.sections] == ["First chapter", "Second chapter"]


def test_chunking_preserves_location_and_offsets() -> None:
    parsed = parse_document(
        filename="notes.txt",
        media_type="text/plain",
        content=("alpha beta gamma delta " * 80).encode(),
        max_extracted_chars=20_000,
    )
    chunks = chunk_document(parsed, chunk_chars=400, overlap_chars=40)

    assert len(chunks) > 1
    assert all(chunk.location == "document" for chunk in chunks)
    assert all(chunk.character_end > chunk.character_start for chunk in chunks)
    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
