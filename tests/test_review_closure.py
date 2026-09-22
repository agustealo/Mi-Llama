from __future__ import annotations

import asyncio
from collections.abc import AsyncIterator, Sequence
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document

from mi_llama.conversations import ConversationService
from mi_llama.documents import _resolve_epub_member, parse_document
from mi_llama.domain import ChatMessage, Conversation, Role, StoredMessage
from mi_llama.providers.errors import ProviderProtocolError, ProviderUnavailableError
from mi_llama.providers.ollama import OllamaProvider
from mi_llama.research import ResearchError, _result_rows
from mi_llama.research_structure.models import (
    CitationCandidate,
    CitationStatus,
    ResearchQuestion,
    ResearchQuestionStatus,
    UpdateResearchQuestionRequest,
)
from mi_llama.research_structure.service import ResearchStructureService
from mi_llama.writing_intelligence.models import AnalyzeManuscriptRequest
from mi_llama.writing_intelligence.service import (
    WritingIntelligenceError,
    WritingIntelligenceService,
    WritingIntelligenceValidationError,
    _sentence_spans,
)
from mi_llama.writing_structure.models import (
    ManuscriptDocument,
    ManuscriptRevision,
    ManuscriptStatus,
)

USER_ID = UUID("11111111-1111-1111-1111-111111111111")
PROJECT_ID = UUID("22222222-2222-2222-2222-222222222222")
CONVERSATION_ID = UUID("33333333-3333-3333-3333-333333333333")
DOCUMENT_ID = UUID("44444444-4444-4444-4444-444444444444")
REVISION_ID = UUID("55555555-5555-5555-5555-555555555555")
CITATION_ID = UUID("66666666-6666-6666-6666-666666666666")
CLAIM_ID = UUID("77777777-7777-7777-7777-777777777777")
EVIDENCE_ID = UUID("88888888-8888-8888-8888-888888888888")
QUESTION_ID = UUID("99999999-9999-9999-9999-999999999999")


def _now() -> datetime:
    return datetime.now(UTC)


def _ollama_provider(transport: httpx.MockTransport) -> OllamaProvider:
    return OllamaProvider(
        base_url="http://ollama.test",
        request_timeout_seconds=5,
        connect_timeout_seconds=1,
        transport=transport,
    )


@pytest.mark.asyncio
async def test_ollama_model_list_rejects_invalid_json() -> None:
    provider = _ollama_provider(
        httpx.MockTransport(lambda _: httpx.Response(200, content=b"not-json"))
    )
    try:
        with pytest.raises(ProviderProtocolError, match="invalid JSON"):
            await provider.list_models()
    finally:
        await provider.close()


@pytest.mark.asyncio
async def test_ollama_stream_requires_terminal_done_event() -> None:
    def transport(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/api/chat"
        return httpx.Response(
            200,
            content=b'{"message":{"role":"assistant","content":"partial"}}\n',
        )

    provider = _ollama_provider(httpx.MockTransport(transport))
    try:
        with pytest.raises(ProviderProtocolError, match="terminal done event"):
            async for _ in provider.chat_stream(
                model="llama3.2:latest",
                messages=[ChatMessage(role=Role.USER, content="hello")],
            ):
                pass
    finally:
        await provider.close()


def test_mindsdb_http_200_error_envelope_is_not_empty_success() -> None:
    with pytest.raises(ResearchError, match="permission denied"):
        _result_rows({"type": "error", "error_message": "permission denied"})


def test_epub_manifest_href_is_uri_resolved_and_fragment_free() -> None:
    resolved = _resolve_epub_member(
        "OPS/package.opf",
        "../Text/chapter%201.xhtml#section-2",
    )
    assert resolved == "Text/chapter 1.xhtml"


def test_docx_extraction_preserves_paragraph_table_paragraph_order() -> None:
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Alpha"
    table.cell(0, 1).text = "Beta"
    document.add_paragraph("After table")
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document(
        filename="ordered.docx",
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=buffer.getvalue(),
        max_extracted_chars=10_000,
    )

    assert parsed.sections[0].text == "Before table\n\nAlpha | Beta\n\nAfter table"


def test_sentence_segmentation_preserves_decimals_and_closing_quotes() -> None:
    spans = _sentence_spans(
        'Revenue rose 3.14 percent. “That matters.” Next claim.',
        base_offset=10,
    )
    assert [span.text for span in spans] == [
        "Revenue rose 3.14 percent.",
        "“That matters.”",
        "Next claim.",
    ]
    assert spans[0].start == 10
    assert spans[1].text.endswith(".”")


class _AnalysisRepository:
    def __init__(self, content: str) -> None:
        now = _now()
        self.document = ManuscriptDocument(
            id=DOCUMENT_ID,
            project_id=PROJECT_ID,
            outline_node_id=None,
            created_by=USER_ID,
            title="Chapter",
            status=ManuscriptStatus.DRAFTING,
            current_revision_id=REVISION_ID,
            current_word_count=1,
            created_at=now,
            updated_at=now,
        )
        self.revision = ManuscriptRevision(
            id=REVISION_ID,
            document_id=DOCUMENT_ID,
            project_id=PROJECT_ID,
            revision_number=1,
            created_by=USER_ID,
            content=content,
            word_count=1,
            created_at=now,
        )

    async def get_manuscript_document(
        self, *, access_token: str, project_id: UUID, document_id: UUID
    ) -> ManuscriptDocument | None:
        del access_token
        if project_id == PROJECT_ID and document_id == DOCUMENT_ID:
            return self.document
        return None

    async def get_manuscript_revision(
        self,
        *,
        access_token: str,
        project_id: UUID,
        document_id: UUID,
        revision_id: UUID,
    ) -> ManuscriptRevision | None:
        del access_token
        if (
            project_id == PROJECT_ID
            and document_id == DOCUMENT_ID
            and revision_id == REVISION_ID
        ):
            return self.revision
        return None


class _FailingStructuredProvider:
    calls = 0

    async def chat_json(
        self,
        *,
        model: str,
        messages: Sequence[ChatMessage],
        schema: dict[str, Any],
    ) -> dict[str, Any]:
        del model, messages, schema
        self.calls += 1
        raise ProviderUnavailableError("offline")


class _UnusedResearch:
    async def query(self, **_: Any) -> Any:
        raise AssertionError("research should not be reached")


@pytest.mark.asyncio
async def test_writing_analysis_rejects_oversized_sentence_before_model_call() -> None:
    provider = _FailingStructuredProvider()
    service = WritingIntelligenceService(
        repository=_AnalysisRepository("x" * 8001),  # type: ignore[arg-type]
        provider=provider,
        research=_UnusedResearch(),  # type: ignore[arg-type]
    )

    with pytest.raises(WritingIntelligenceValidationError, match="at most 8000"):
        await service.analyze(
            access_token="jwt",
            project_id=PROJECT_ID,
            document_id=DOCUMENT_ID,
            request=AnalyzeManuscriptRequest(model="llama3.2:latest"),
        )
    assert provider.calls == 0


@pytest.mark.asyncio
async def test_writing_analysis_wraps_structured_provider_failure() -> None:
    service = WritingIntelligenceService(
        repository=_AnalysisRepository("A factual claim exists."),  # type: ignore[arg-type]
        provider=_FailingStructuredProvider(),
        research=_UnusedResearch(),  # type: ignore[arg-type]
    )

    with pytest.raises(WritingIntelligenceError, match="claim selection failed"):
        await service.analyze(
            access_token="jwt",
            project_id=PROJECT_ID,
            document_id=DOCUMENT_ID,
            request=AnalyzeManuscriptRequest(model="llama3.2:latest"),
        )


class _ResearchReviewRepository:
    def __init__(self) -> None:
        now = _now()
        self.question = ResearchQuestion(
            id=QUESTION_ID,
            project_id=PROJECT_ID,
            created_by=USER_ID,
            question="What happened?",
            status=ResearchQuestionStatus.RESOLVED,
            priority=3,
            resolution="Existing resolution",
            created_at=now,
            updated_at=now,
        )
        self.citation = CitationCandidate(
            id=CITATION_ID,
            project_id=PROJECT_ID,
            claim_id=CLAIM_ID,
            evidence_id=EVIDENCE_ID,
            created_by=USER_ID,
            status=CitationStatus.ACCEPTED,
            created_at=now,
            updated_at=now,
        )
        self.signal_count = 0
        self.citation_update_count = 0

    async def get_research_question(self, **_: Any) -> ResearchQuestion:
        return self.question

    async def update_research_question(self, **_: Any) -> ResearchQuestion:
        raise AssertionError("invalid resolved update must not reach persistence")

    async def get_citation_candidate(self, **_: Any) -> CitationCandidate:
        return self.citation

    async def update_citation_candidate(self, **_: Any) -> CitationCandidate:
        self.citation_update_count += 1
        return self.citation

    async def add_learning_signal(self, **_: Any) -> None:
        self.signal_count += 1


@pytest.mark.asyncio
async def test_resolved_question_rejects_explicit_null_resolution() -> None:
    service = ResearchStructureService(
        repository=_ResearchReviewRepository(),  # type: ignore[arg-type]
        research=None,
    )
    with pytest.raises(ValueError, match="require a resolution"):
        await service.update_question(
            access_token="jwt",
            project_id=PROJECT_ID,
            question_id=QUESTION_ID,
            request=UpdateResearchQuestionRequest(resolution=None),
        )


@pytest.mark.asyncio
async def test_repeated_citation_status_is_idempotent() -> None:
    repository = _ResearchReviewRepository()
    service = ResearchStructureService(
        repository=repository,  # type: ignore[arg-type]
        research=None,
    )
    result = await service.update_citation(
        access_token="jwt",
        project_id=PROJECT_ID,
        citation_id=CITATION_ID,
        citation_status=CitationStatus.ACCEPTED,
    )
    assert result.status is CitationStatus.ACCEPTED
    assert repository.citation_update_count == 0
    assert repository.signal_count == 0


class _ConversationRepository:
    def __init__(self) -> None:
        now = _now()
        self.conversation = Conversation(
            id=CONVERSATION_ID,
            project_id=PROJECT_ID,
            created_by=USER_ID,
            title="Thread",
            model="llama3.2:latest",
            created_at=now,
            updated_at=now,
        )
        self.messages: list[StoredMessage] = []

    async def get_conversation(self, **_: Any) -> Conversation:
        return self.conversation

    async def add_message(
        self,
        *,
        conversation_id: UUID,
        role: Role,
        content: str,
        **_: Any,
    ) -> StoredMessage:
        message = StoredMessage(
            id=uuid4(),
            conversation_id=conversation_id,
            created_by=USER_ID,
            role=role,
            content=content,
            created_at=_now(),
        )
        self.messages.append(message)
        return message

    async def get_messages(self, **_: Any) -> list[StoredMessage]:
        return list(self.messages)


class _BlockingChatProvider:
    def __init__(self) -> None:
        self.first_entered = asyncio.Event()
        self.release_first = asyncio.Event()
        self.calls = 0

    async def chat_stream(
        self,
        *,
        model: str,
        messages: Sequence[ChatMessage],
    ) -> AsyncIterator[str]:
        del model
        self.calls += 1
        if self.calls == 1:
            self.first_entered.set()
            await self.release_first.wait()
        yield f"reply:{messages[-1].content}"


@pytest.mark.asyncio
async def test_conversation_turns_do_not_interleave() -> None:
    repository = _ConversationRepository()
    provider = _BlockingChatProvider()
    service = ConversationService(
        repository=repository,  # type: ignore[arg-type]
        provider=provider,  # type: ignore[arg-type]
    )

    async def consume(content: str) -> str:
        chunks = [
            chunk
            async for chunk in service.stream_reply(
                access_token="jwt",
                conversation_id=CONVERSATION_ID,
                user_content=content,
            )
        ]
        return "".join(chunks)

    first = asyncio.create_task(consume("first"))
    await provider.first_entered.wait()
    second = asyncio.create_task(consume("second"))
    await asyncio.sleep(0)
    assert [(message.role, message.content) for message in repository.messages] == [
        (Role.USER, "first")
    ]

    provider.release_first.set()
    assert await first == "reply:first"
    assert await second == "reply:second"
    assert [(message.role, message.content) for message in repository.messages] == [
        (Role.USER, "first"),
        (Role.ASSISTANT, "reply:first"),
        (Role.USER, "second"),
        (Role.ASSISTANT, "reply:second"),
    ]


def test_review_closure_migration_contains_durable_invariants() -> None:
    sql = Path("supabase/migrations/20260922014500_review_closure_hardening.sql").read_text(
        encoding="utf-8"
    )

    assert "drop constraint if exists conversations_created_by_fkey" in sql
    assert "drop constraint if exists writing_analysis_findings_created_by_fkey" in sql
    assert "messages_touch_conversation_activity" in sql
    assert "acquire_conversation_reply_lease" in sql
    assert "conversation reply already in progress" in sql
    assert "pg_advisory_xact_lock" in sql
    assert "citation is linked to manuscript writing and must remain accepted" in sql
    assert "create_manuscript_revision_result" in sql
    assert "to_jsonb(updated_document), to_jsonb(inserted_revision)" in sql
