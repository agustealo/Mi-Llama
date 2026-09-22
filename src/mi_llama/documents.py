from __future__ import annotations

import posixpath
import re
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path, PurePosixPath
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from mi_llama.domain import SourceKind


class DocumentExtractionError(ValueError):
    """The uploaded document cannot be safely parsed into research text."""


@dataclass(frozen=True)
class ParsedSection:
    location: str | None
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    kind: SourceKind
    parser: str
    sections: tuple[ParsedSection, ...]

    @property
    def character_count(self) -> int:
        return sum(len(section.text) for section in self.sections)


@dataclass(frozen=True)
class PreparedChunk:
    ordinal: int
    location: str | None
    content: str
    character_start: int
    character_end: int


class _ReadableHTMLParser(HTMLParser):
    _BLOCK_TAGS = {
        "article",
        "blockquote",
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "p",
        "section",
        "td",
        "th",
        "tr",
    }
    _SKIP_TAGS = {"script", "style", "svg", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        lowered = tag.lower()
        if lowered in self._SKIP_TAGS:
            self._skip_depth += 1
        elif self._skip_depth == 0 and lowered in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        lowered = tag.lower()
        if lowered in self._SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif self._skip_depth == 0 and lowered in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            self._parts.append(data)

    def text(self) -> str:
        return _normalize_text("".join(self._parts))


def detect_source_kind(filename: str, media_type: str) -> SourceKind:
    suffix = Path(filename).suffix.lower()
    media = media_type.lower().split(";", 1)[0].strip()
    if suffix == ".pdf" or media == "application/pdf":
        return SourceKind.PDF
    if suffix == ".docx" or media == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return SourceKind.DOCX
    if suffix == ".epub" or media == "application/epub+zip":
        return SourceKind.EPUB
    if suffix in {".md", ".markdown"} or media in {"text/markdown", "text/x-markdown"}:
        return SourceKind.MARKDOWN
    if suffix in {".html", ".htm"} or media in {"text/html", "application/xhtml+xml"}:
        return SourceKind.HTML
    if suffix in {".txt", ".text"} or media.startswith("text/plain"):
        return SourceKind.TEXT
    raise DocumentExtractionError(
        "Unsupported source type. Supported formats are PDF, DOCX, EPUB, TXT, Markdown, and HTML."
    )


def parse_document(
    *,
    filename: str,
    media_type: str,
    content: bytes,
    max_extracted_chars: int,
) -> ParsedDocument:
    kind = detect_source_kind(filename, media_type)
    if not content:
        raise DocumentExtractionError("The uploaded source is empty")

    if kind is SourceKind.PDF:
        document = _parse_pdf(content)
    elif kind is SourceKind.DOCX:
        document = _parse_docx(content)
    elif kind is SourceKind.EPUB:
        document = _parse_epub(content)
    elif kind is SourceKind.HTML:
        document = _parse_html(content)
    elif kind is SourceKind.MARKDOWN:
        document = _parse_text(content, kind=SourceKind.MARKDOWN, parser="markdown-text-v1")
    else:
        document = _parse_text(content, kind=SourceKind.TEXT, parser="plain-text-v1")

    if not document.sections or not any(section.text.strip() for section in document.sections):
        raise DocumentExtractionError("No readable text could be extracted from the source")
    if document.character_count > max_extracted_chars:
        raise DocumentExtractionError(
            f"Extracted source exceeds the {max_extracted_chars} character safety limit"
        )
    return document


def chunk_document(
    document: ParsedDocument,
    *,
    chunk_chars: int,
    overlap_chars: int,
) -> list[PreparedChunk]:
    if overlap_chars >= chunk_chars:
        raise ValueError("Chunk overlap must be smaller than chunk size")
    chunks: list[PreparedChunk] = []
    global_offset = 0
    ordinal = 0
    for section in document.sections:
        text = _normalize_text(section.text)
        if not text:
            continue
        local_start = 0
        while local_start < len(text):
            target_end = min(len(text), local_start + chunk_chars)
            local_end = _word_boundary_end(text, local_start, target_end)
            chunk_text = text[local_start:local_end].strip()
            if chunk_text:
                chunks.append(
                    PreparedChunk(
                        ordinal=ordinal,
                        location=section.location,
                        content=chunk_text,
                        character_start=global_offset + local_start,
                        character_end=global_offset + local_end,
                    )
                )
                ordinal += 1
            if local_end >= len(text):
                break
            next_start = max(local_start + 1, local_end - overlap_chars)
            while next_start < local_end and not text[next_start].isspace():
                next_start += 1
            local_start = min(next_start, local_end)
        global_offset += len(text) + 1
    if not chunks:
        raise DocumentExtractionError("The source did not produce any research chunks")
    return chunks


def _parse_pdf(content: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            result = reader.decrypt("")
            if result == 0:
                raise DocumentExtractionError("Encrypted PDFs require removal of the password")
        sections = tuple(
            ParsedSection(location=f"page {index}", text=_normalize_text(page.extract_text() or ""))
            for index, page in enumerate(reader.pages, start=1)
        )
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError("PDF parsing failed") from exc
    return ParsedDocument(kind=SourceKind.PDF, parser="pypdf-v1", sections=sections)


def _parse_docx(content: bytes) -> ParsedDocument:
    try:
        document = Document(BytesIO(content))
        parts: list[str] = []
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                text = Paragraph(child, document).text.strip()
                if text:
                    parts.append(text)
            elif child.tag.endswith("}tbl"):
                table = Table(child, document)
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
    except Exception as exc:
        raise DocumentExtractionError("DOCX parsing failed") from exc
    text = _normalize_text("\n\n".join(parts))
    return ParsedDocument(
        kind=SourceKind.DOCX,
        parser="python-docx-v2",
        sections=(ParsedSection(location="document", text=text),),
    )


def _parse_epub(content: bytes) -> ParsedDocument:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            ordered = _epub_spine_documents(archive)
            if not ordered:
                ordered = sorted(
                    name
                    for name in archive.namelist()
                    if name.lower().endswith((".xhtml", ".html", ".htm"))
                )
            sections: list[ParsedSection] = []
            for index, name in enumerate(ordered, start=1):
                parser = _ReadableHTMLParser()
                parser.feed(_decode_text(archive.read(name)))
                text = parser.text()
                if text:
                    sections.append(
                        ParsedSection(
                            location=f"section {index}: {PurePosixPath(name).name}",
                            text=text,
                        )
                    )
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise DocumentExtractionError("EPUB parsing failed") from exc
    return ParsedDocument(kind=SourceKind.EPUB, parser="epub-zip-v2", sections=tuple(sections))


def _epub_spine_documents(archive: zipfile.ZipFile) -> list[str]:
    container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
    rootfile = next(
        (
            node.attrib.get("full-path")
            for node in container.iter()
            if node.tag.endswith("rootfile") and node.attrib.get("full-path")
        ),
        None,
    )
    if rootfile is None:
        return []
    package = ElementTree.fromstring(archive.read(rootfile))
    manifest: dict[str, str] = {}
    for node in package.iter():
        if node.tag.endswith("item"):
            item_id = node.attrib.get("id")
            href = node.attrib.get("href")
            if item_id and href:
                manifest[item_id] = href
    ordered: list[str] = []
    for node in package.iter():
        if node.tag.endswith("itemref"):
            href = manifest.get(node.attrib.get("idref", ""))
            if href:
                ordered.append(_resolve_epub_member(rootfile, href))
    return ordered


def _resolve_epub_member(rootfile: str, href: str) -> str:
    parsed = urlsplit(href)
    if parsed.scheme or parsed.netloc:
        raise DocumentExtractionError("EPUB manifest contains an external document reference")
    decoded_path = unquote(parsed.path)
    if not decoded_path:
        raise DocumentExtractionError("EPUB manifest contains an empty document reference")
    if decoded_path.startswith("/"):
        raise DocumentExtractionError("EPUB manifest document reference escapes the archive root")

    base = posixpath.dirname(rootfile)
    resolved = posixpath.normpath(posixpath.join(base, decoded_path))
    if resolved == ".." or resolved.startswith("../") or resolved.startswith("/"):
        raise DocumentExtractionError("EPUB manifest document reference escapes the archive root")
    return resolved


def _parse_html(content: bytes) -> ParsedDocument:
    parser = _ReadableHTMLParser()
    try:
        parser.feed(_decode_text(content))
    except Exception as exc:
        raise DocumentExtractionError("HTML parsing failed") from exc
    return ParsedDocument(
        kind=SourceKind.HTML,
        parser="html-parser-v1",
        sections=(ParsedSection(location="document", text=parser.text()),),
    )


def _parse_text(content: bytes, *, kind: SourceKind, parser: str) -> ParsedDocument:
    return ParsedDocument(
        kind=kind,
        parser=parser,
        sections=(ParsedSection(location="document", text=_normalize_text(_decode_text(content))),),
    )


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return content.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError("Text encoding is not supported") from exc
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252")


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\f\v ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _word_boundary_end(text: str, start: int, target_end: int) -> int:
    if target_end >= len(text):
        return len(text)
    minimum = start + max(1, int((target_end - start) * 0.7))
    for index in range(target_end, minimum, -1):
        if text[index - 1].isspace():
            return index
    return target_end
